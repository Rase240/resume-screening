import csv
import hashlib
import io
import json
import os
import sqlite3
import uuid
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader

from scoring import parse_resume, score

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
FILES = DATA / "files"
FILES.mkdir(parents=True, exist_ok=True)
DB = DATA / "folio.db"
MAX_FILE = 10 * 1024 * 1024
WEIGHT_KEYS = {"Required skills", "Preferred skills", "Experience", "Education", "Location", "Language similarity"}
app = FastAPI(title="Folio screening API", version="1.0")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

def conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys = ON")
    return c

def init():
    with conn() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS roles(id TEXT PRIMARY KEY, title TEXT NOT NULL, description TEXT NOT NULL, requirements TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS candidates(id TEXT PRIMARY KEY, role_id TEXT NOT NULL REFERENCES roles(id) ON DELETE CASCADE, filename TEXT NOT NULL, sha256 TEXT NOT NULL, file_path TEXT NOT NULL, text TEXT NOT NULL, parsed TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'New', notes TEXT NOT NULL DEFAULT '', created_at TEXT DEFAULT CURRENT_TIMESTAMP, UNIQUE(role_id, sha256));
        """)
init()

class RoleIn(BaseModel):
    title: str = Field(min_length=2, max_length=120)
    description: str = Field(min_length=10, max_length=20000)
    required_skills: list[str] = []
    preferred_skills: list[str] = []
    min_experience: float = Field(default=0, ge=0, le=50)
    education: str = ""
    location: str = ""
    work_mode: str = "Hybrid"
    weights: dict[str, int] = {"Required skills":35,"Preferred skills":10,"Experience":15,"Education":10,"Location":5,"Language similarity":25}

class CandidateEdit(BaseModel):
    status: str | None = None
    notes: str | None = None
    text: str | None = None
    parsed: dict | None = None

class BulkStatusEdit(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=50)
    status: str

def role_from(row):
    return {"id": row["id"], **json.loads(row["requirements"]), "title": row["title"], "description": row["description"], "created_at": row["created_at"]}

def get_role(role_id):
    with conn() as c:
        row = c.execute("SELECT * FROM roles WHERE id=?", (role_id,)).fetchone()
    if not row: raise HTTPException(404, "Role not found")
    return role_from(row)

def candidate_from(row, role):
    item = {k: row[k] for k in ["id", "role_id", "filename", "text", "status", "notes", "created_at"]}
    item["parsed"] = json.loads(row["parsed"])
    item["match"] = score(role, item)
    return item

@app.get("/health")
def health(): return {"status": "ok"}

@app.get("/roles")
def list_roles():
    with conn() as c: rows = c.execute("SELECT * FROM roles ORDER BY created_at DESC").fetchall()
    return [role_from(r) for r in rows]

@app.post("/roles", status_code=201)
def create_role(body: RoleIn):
    if set(body.weights) != WEIGHT_KEYS or sum(body.weights.values()) != 100 or any(v < 0 for v in body.weights.values()):
        raise HTTPException(422, "Scoring weights must total 100")
    role_id = str(uuid.uuid4())
    info = body.model_dump()
    with conn() as c:
        c.execute("INSERT INTO roles(id,title,description,requirements) VALUES(?,?,?,?)", (role_id, body.title, body.description, json.dumps(info)))
    return get_role(role_id)

@app.put("/roles/{role_id}")
def update_role(role_id: str, body: RoleIn):
    get_role(role_id)
    if set(body.weights) != WEIGHT_KEYS or sum(body.weights.values()) != 100 or any(v < 0 for v in body.weights.values()): raise HTTPException(422, "Scoring weights must total 100")
    with conn() as c: c.execute("UPDATE roles SET title=?,description=?,requirements=? WHERE id=?", (body.title, body.description, json.dumps(body.model_dump()), role_id))
    return get_role(role_id)

@app.delete("/roles/{role_id}", status_code=204)
def delete_role(role_id: str):
    get_role(role_id)
    with conn() as c:
        paths = [r[0] for r in c.execute("SELECT file_path FROM candidates WHERE role_id=?",(role_id,))]
        c.execute("DELETE FROM roles WHERE id=?", (role_id,))
    for path in paths: Path(path).unlink(missing_ok=True)

@app.get("/roles/{role_id}/candidates")
def list_candidates(role_id: str):
    role = get_role(role_id)
    with conn() as c: rows = c.execute("SELECT * FROM candidates WHERE role_id=?", (role_id,)).fetchall()
    return sorted([candidate_from(r, role) for r in rows], key=lambda x: (-x["match"]["score"], x["parsed"]["name"].lower()))

def extract(data, suffix):
    stream = io.BytesIO(data)
    if suffix == ".pdf":
        reader = PdfReader(stream)
        if len(reader.pages) > 40: raise ValueError("PDF exceeds 40 pages")
        return "\n".join(p.extract_text() or "" for p in reader.pages)
    doc = Document(stream)
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)

@app.post("/roles/{role_id}/upload")
async def upload(role_id: str, files: list[UploadFile] = File(...)):
    get_role(role_id)
    if len(files) > 50: raise HTTPException(413, "Upload at most 50 files at once")
    results = []
    for file in files:
        name = Path(file.filename or "resume").name[:150]
        suffix = Path(name).suffix.lower()
        if suffix not in [".pdf", ".docx"]:
            results.append({"filename": name, "status": "error", "message": "Only PDF and DOCX files are supported"}); continue
        data = await file.read(MAX_FILE + 1)
        if len(data) > MAX_FILE:
            results.append({"filename":name,"status":"error","message":"File exceeds 10 MB"}); continue
        if not ((suffix == ".pdf" and data.startswith(b"%PDF-")) or (suffix == ".docx" and data.startswith(b"PK"))):
            results.append({"filename":name,"status":"error","message":"File content does not match extension"}); continue
        digest = hashlib.sha256(data).hexdigest()
        with conn() as c: duplicate = c.execute("SELECT id FROM candidates WHERE role_id=? AND sha256=?", (role_id, digest)).fetchone()
        if duplicate:
            results.append({"filename":name,"status":"duplicate","message":"Already uploaded for this role"}); continue
        try:
            text = extract(data, suffix).strip()
            if len(text) < 40: raise ValueError("No readable text found. Scanned PDFs need OCR.")
            parsed = parse_resume(text, name)
            candidate_id = str(uuid.uuid4())
            path = FILES / (candidate_id + suffix)
            path.write_bytes(data)
            with conn() as c:
                c.execute("INSERT INTO candidates(id,role_id,filename,sha256,file_path,text,parsed) VALUES(?,?,?,?,?,?,?)", (candidate_id,role_id,name,digest,str(path),text,json.dumps(parsed)))
            results.append({"filename":name,"status":"processed","id":candidate_id,"name":parsed["name"]})
        except Exception as exc:
            results.append({"filename":name,"status":"error","message":str(exc)[:180]})
    return results

def get_candidate(candidate_id):
    with conn() as c: row = c.execute("SELECT * FROM candidates WHERE id=?", (candidate_id,)).fetchone()
    if not row: raise HTTPException(404, "Candidate not found")
    return row

@app.patch("/candidates/{candidate_id}")
def edit_candidate(candidate_id: str, body: CandidateEdit):
    row = get_candidate(candidate_id)
    status = body.status if body.status is not None else row["status"]
    if status not in ["New", "In review", "Shortlisted", "Rejected"]: raise HTTPException(422, "Invalid status")
    notes = body.notes if body.notes is not None else row["notes"]
    text = body.text if body.text is not None else row["text"]
    parsed = body.parsed if body.parsed is not None else json.loads(row["parsed"])
    if len(text) > 100000 or len(notes) > 5000: raise HTTPException(422, "Text too long")
    with conn() as c: c.execute("UPDATE candidates SET status=?,notes=?,text=?,parsed=? WHERE id=?", (status,notes,text,json.dumps(parsed),candidate_id))
    return candidate_from(get_candidate(candidate_id), get_role(row["role_id"]))

@app.patch("/roles/{role_id}/candidates/status")
def bulk_status(role_id: str, body: BulkStatusEdit):
    get_role(role_id)
    if body.status not in ["New", "In review", "Shortlisted", "Rejected"]:
        raise HTTPException(422, "Invalid status")
    ids = list(dict.fromkeys(body.ids))
    with conn() as c:
        rows = c.execute(
            f"SELECT id FROM candidates WHERE role_id=? AND id IN ({','.join('?' for _ in ids)})",
            (role_id, *ids),
        ).fetchall()
        if len(rows) != len(ids):
            raise HTTPException(404, "One or more candidates were not found in this role")
        c.executemany("UPDATE candidates SET status=? WHERE id=?", [(body.status, candidate_id) for candidate_id in ids])
    return {"updated": len(ids), "status": body.status}
