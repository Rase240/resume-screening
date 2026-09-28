import io
from docx import Document
from fastapi.testclient import TestClient
from main import app

client=TestClient(app)
ROLE={"title":"Test Data Engineer","description":"Build data pipelines in Python and SQL", "required_skills":["Python","SQL"],"preferred_skills":["AWS"],"min_experience":2,"education":"","location":"","work_mode":"Remote"}

def test_full_role_upload_review_export_delete_flow():
    role=client.post('/roles',json=ROLE).json()
    rid=role['id']
    doc=Document();doc.add_paragraph('Casey Sample');doc.add_paragraph('Data Engineer');doc.add_paragraph('casey@example.test');doc.add_paragraph('4 years using Python, SQL, and AWS to build data pipelines.')
    raw=io.BytesIO();doc.save(raw)
    files={'files':('casey.docx',raw.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    response=client.post(f'/roles/{rid}/upload',files=files)
    assert response.status_code==200 and response.json()[0]['status']=='processed'
    assert client.post(f'/roles/{rid}/upload',files=files).json()[0]['status']=='duplicate'
    candidate=client.get(f'/roles/{rid}/candidates').json()[0]
    assert candidate['match']['score']>0
    assert client.patch(f'/roles/{rid}/candidates/status',json={'ids':[candidate['id']],'status':'In review'}).json()=={'updated':1,'status':'In review'}
    assert client.get(f'/roles/{rid}/candidates').json()[0]['status']=='In review'
    assert client.patch(f'/roles/{rid}/candidates/status',json={'ids':['missing'],'status':'Rejected'}).status_code==404
    assert client.patch(f"/candidates/{candidate['id']}",json={'status':'Shortlisted','notes':'Review portfolio'}).json()['status']=='Shortlisted'
    assert 'Casey Sample' in client.get(f'/roles/{rid}/export').text
    assert client.get(f"/candidates/{candidate['id']}/file").status_code==200
    assert client.delete(f"/candidates/{candidate['id']}").status_code==204
    assert client.get(f'/roles/{rid}/candidates').json()==[]
    assert client.delete(f'/roles/{rid}').status_code==204
